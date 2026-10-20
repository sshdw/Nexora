# -*- coding: utf-8 -*-
import re, copy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from content import BLOCKS, TITLE, SUBTITLE, META

ACCENT = RGBColor(0x0F, 0x5C, 0x4A)       # deep green
ACCENT2 = RGBColor(0x1B, 0x6E, 0x8C)      # teal
DARK = RGBColor(0x1F, 0x2A, 0x2E)
GREY = RGBColor(0x55, 0x60, 0x66)
ORANGE = RGBColor(0xB5, 0x5A, 0x00)

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


def shade(el, fill):
    tcPr = el.get_or_add_tcPr() if el.tag.endswith('}tc') else el
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def para_shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def para_border(p, color="9BBF8A", size=18, left_only=True, space=6):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    sides = ['left'] if left_only else ['top', 'left', 'bottom', 'right']
    for side in sides:
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(size))
        e.set(qn('w:space'), str(space))
        e.set(qn('w:color'), color)
        bdr.append(e)
    pPr.append(bdr)


def cell_margins(cell, top=90, bottom=90, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for name, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + name)
        e.set(qn('w:w'), str(val))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tcPr.append(mar)


def table_borders(table, color="C9D6CF", size=6):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(size))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        borders.append(e)
    tblPr.append(borders)


BOLD_RE = re.compile(r'\*\*(.+?)\*\*')


def add_runs(p, text, size=10.5, color=DARK, font=BODY_FONT, bold=False):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            r.font.size, r.font.color.rgb, r.font.name, r.bold = Pt(size), color, font, bold
        r = p.add_run(m.group(1))
        r.font.size, r.font.color.rgb, r.font.name, r.bold = Pt(size), ACCENT, font, True
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        r.font.size, r.font.color.rgb, r.font.name, r.bold = Pt(size), color, font, bold


def set_spacing(p, before=0, after=4, line=1.12):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Химия · 9-сынып конспектісі · ")
    run.font.size, run.font.color.rgb = Pt(8), GREY
    run2 = p.add_run()
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run2._r.append(fld1); run2._r.append(instr); run2._r.append(fld2)
    run2.font.size, run2.font.color.rgb = Pt(8), GREY


def build(path):
    doc = Document()

    st = doc.styles['Normal']
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
    st.paragraph_format.space_after = Pt(4)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.6)

    add_page_number(sec)

    # ---------------- COVER ----------------
    p = doc.add_paragraph(); set_spacing(p, 30, 0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(p, 0, 2)
    r = p.add_run("ХИМИЯ")
    r.font.size, r.bold, r.font.color.rgb = Pt(46), True, ACCENT
    r.font.name = BODY_FONT
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(p, 0, 18)
    r = p.add_run("9-СЫНЫП · КОНСПЕКТ")
    r.font.size, r.bold, r.font.color.rgb = Pt(20), True, ACCENT2

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(p, 0, 6)
    r = p.add_run(SUBTITLE)
    r.font.size, r.bold, r.font.color.rgb = Pt(17), True, DARK

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(p, 0, 30)
    r = p.add_run("Оксидтер · Негіздер · Қышқылдар · Тұздар")
    r.font.size, r.font.color.rgb = Pt(12), ORANGE

    for line in META:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_spacing(p, 0, 3)
        r = p.add_run(line)
        r.font.size, r.font.color.rgb = Pt(10), GREY

    # cover quick-map table
    p = doc.add_paragraph(); set_spacing(p, 26, 8)
    r = p.add_run("КОНСПЕКТ ҚҰРЫЛЫМЫ")
    r.font.size, r.bold, r.font.color.rgb = Pt(11), True, ACCENT

    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, "D6E2DA", 6)
    t.autofit = False
    rows = [
        ("1", "Периодтық жүйе заңдылықтары"),
        ("2", "Оксидтер"),
        ("3", "Негіздер"),
        ("4", "Қышқылдар"),
        ("5", "Тұздар"),
        ("6–7", "Генетикалық байланыс · Иондық теңдеулер"),
        ("8", "Зертханалық тәжірибелер"),
        ("9–10", "Жиі жіберілетін қателер · Мини-тест"),
    ]
    for num, name in rows:
        row = t.add_row()
        row.cells[0].width = Cm(2.2)
        row.cells[1].width = Cm(14.6)
        for c in row.cells:
            cell_margins(c, 60, 60, 120, 120)
        pr = row.cells[0].paragraphs[0]
        rr = pr.add_run(num); rr.bold = True; rr.font.size = Pt(10); rr.font.color.rgb = ACCENT
        pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pr2 = row.cells[1].paragraphs[0]
        rr = pr2.add_run(name); rr.font.size = Pt(10.5)

    p = doc.add_paragraph(); set_spacing(p, 30, 0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Негізгі мазмұн НИШ 9-сынып бағдарламасының 9.1В және 9.4 бөлімдері бойынша жинақталған")
    r.font.size, r.italic, r.font.color.rgb = Pt(8.5), True, GREY

    doc.add_page_break()

    # ---------------- BODY ----------------
    for block in BLOCKS:
        kind = block[0]

        if kind == "h1":
            p = doc.add_paragraph()
            set_spacing(p, 14, 7, 1.0)
            para_shade(p, "0F5C4A")
            p.paragraph_format.left_indent = Cm(0.1)
            r = p.add_run("  " + block[1])
            r.font.size, r.bold = Pt(13.5), True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.name = BODY_FONT
            p.paragraph_format.keep_with_next = True

        elif kind == "h2":
            p = doc.add_paragraph()
            set_spacing(p, 10, 4, 1.0)
            r = p.add_run(block[1])
            r.font.size, r.bold, r.font.color.rgb = Pt(11.5), True, ACCENT2
            p.paragraph_format.keep_with_next = True

        elif kind == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_spacing(p, 2, 5, 1.15)
            add_runs(p, block[1], 10.5)

        elif kind == "bullets":
            for item in block[1]:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                set_spacing(p, 0, 3, 1.12)
                pf = p.paragraph_format
                pf.left_indent = Cm(0.65)
                pf.first_line_indent = Cm(-0.4)
                r = p.add_run("▪  ")
                r.font.size, r.font.color.rgb, r.bold = Pt(10), ACCENT, True
                add_runs(p, item, 10.5)

        elif kind == "nums":
            for i, item in enumerate(block[1], 1):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                set_spacing(p, 0, 3, 1.12)
                pf = p.paragraph_format
                pf.left_indent = Cm(0.75)
                pf.first_line_indent = Cm(-0.55)
                r = p.add_run(f"{i}. ")
                r.font.size, r.bold, r.font.color.rgb = Pt(10.5), True, ACCENT
                add_runs(p, item, 10.5)

        elif kind == "note":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_spacing(p, 5, 8, 1.1)
            para_shade(p, "FFF4E2")
            para_border(p, "E0A050", 20)
            p.paragraph_format.left_indent = Cm(0.25)
            r = p.add_run("!  ")
            r.font.size, r.bold, r.font.color.rgb = Pt(10.5), True, ORANGE
            add_runs(p, block[1], 9.5, RGBColor(0x6B, 0x45, 0x10))

        elif kind == "box":
            t = doc.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_borders(t, "B8D3AC", 8)
            cell = t.rows[0].cells[0]
            cell.width = Cm(16.9)
            shade(cell._tc, "F3F8F0")
            cell_margins(cell, 110, 110, 160, 120)
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
            for line in block[1]:
                cp = cell.add_paragraph()
                set_spacing(cp, 0, 1, 1.05)
                if line == "":
                    cp.add_run("").font.size = Pt(4)
                    continue
                r = cp.add_run(line)
                r.font.name = MONO_FONT
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x14, 0x3D, 0x33)
            sp = doc.add_paragraph()
            set_spacing(sp, 0, 2, 1.0)
            sp.add_run("").font.size = Pt(3)

        elif kind == "pagebreak":
            doc.add_page_break()

        elif kind == "table":
            headers, rows, widths = block[1], block[2], block[3]
            t = doc.add_table(rows=1, cols=len(headers))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            table_borders(t, "CBD8D1", 6)
            hdr = t.rows[0]
            for i, h in enumerate(headers):
                c = hdr.cells[i]
                c.width = Cm(widths[i])
                shade(c._tc, "0F5C4A")
                cell_margins(c, 70, 70, 110, 110)
                pr = c.paragraphs[0]
                set_spacing(pr, 0, 0, 1.0)
                r = pr.add_run(h)
                r.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            # header repeat
            trPr = hdr._tr.get_or_add_trPr()
            th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)

            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                for i, val in enumerate(row):
                    c = cells[i]
                    c.width = Cm(widths[i])
                    if ri % 2 == 1:
                        shade(c._tc, "F1F6F3")
                    cell_margins(c, 70, 70, 110, 110)
                    pr = c.paragraphs[0]
                    set_spacing(pr, 0, 0, 1.05)
                    r = pr.add_run(val)
                    r.font.size = Pt(9)
                    r.font.color.rgb = DARK
                    if i == 0:
                        r.bold = True
            sp = doc.add_paragraph()
            set_spacing(sp, 0, 4, 1.0)
            sp.add_run("").font.size = Pt(4)

    doc.save(path)
    print("saved", path)


if __name__ == "__main__":
    build("Химия_9_сынып_конспект.docx")
